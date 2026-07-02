#!/usr/bin/env python3
"""Build a concise DOCX proposal for the ProgramBench scaffold project."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/ProgramBench_Scaffolding_Proposal.docx")


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(85, 85, 85)
LIGHT_GRAY = "F2F4F7"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa: int = 9360) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")


def set_col_widths(table, widths_in: list[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths_in):
            row.cells[idx].width = Inches(width)
            set_cell_margins(row.cells[idx])
            row.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("ProgramBench scaffold proposal")
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = GRAY


def add_title(doc: Document) -> None:
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("Behavioral Specification Scaffolding for Black-Box Program Reconstruction Agents")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = DARK_BLUE

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("ProgramBench proposal: agent-interface design, uncertainty-guided probing, and oracle-grounded repair")
    run.font.size = Pt(11)
    run.font.color.rgb = GRAY

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(12)
    run = meta.add_run("Prepared: June 20, 2026 | Current scope: agent scaffolding, not model training")
    run.font.size = Pt(10)
    run.font.color.rgb = GRAY


def add_key_value_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_width(table)
    set_col_widths(table, [1.8, 4.7])
    hdr = table.rows[0].cells
    hdr[0].text = "Item"
    hdr[1].text = "Summary"
    for cell in hdr:
        set_cell_shading(cell, LIGHT_GRAY)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for label, text in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = text
        for cell in cells:
            set_cell_margins(cell)
    doc.add_paragraph()


def add_bullet(doc: Document, text: str) -> None:
    para = doc.add_paragraph(text, style="List Bullet")
    para.paragraph_format.space_after = Pt(4)


def add_number(doc: Document, text: str) -> None:
    para = doc.add_paragraph(text, style="List Number")
    para.paragraph_format.space_after = Pt(4)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_width(table)
    hdr = table.rows[0].cells
    for idx, header in enumerate(headers):
        hdr[idx].text = header
        set_cell_shading(hdr[idx], LIGHT_GRAY)
        for paragraph in hdr[idx].paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cells[idx].text = text
    set_col_widths(table, widths)
    doc.add_paragraph()


def build_doc() -> None:
    doc = Document()
    style_document(doc)
    add_title(doc)

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(
        "ProgramBench asks agents to rebuild a program from only a compiled executable and documentation. "
        "Our proposal is not to train a new model, but to design a better agent-side interface for this setting: "
        "a behavioral specification scaffold that turns black-box executable observations into structured, "
        "implementation-ready evidence for a coding agent."
    )
    add_key_value_table(
        doc,
        [
            (
                "Core claim",
                "ProgramBench failures are partly specification-recovery failures, not only code-generation failures.",
            ),
            (
                "Proposed contribution",
                "A behavioral specification scaffold that builds a behavior map, selects informative probes, and returns oracle-grounded counterexamples.",
            ),
            (
                "MVP",
                "Compare the same baseline agent with and without the scaffold on 5-10 easy ProgramBench CLI tasks.",
            ),
            (
                "Not current scope",
                "SFT, RL, reward-model training, or full-benchmark scaling. These are future work.",
            ),
        ],
    )

    doc.add_heading("Problem Setting", level=1)
    doc.add_paragraph(
        "In ProgramBench, the agent sees a reference executable and task documentation, but not the original source code. "
        "It must produce a fresh source codebase whose built executable matches the reference behavior under hidden tests."
    )
    add_bullet(doc, "Inputs available to the agent: documentation, reference executable, and local command execution.")
    add_bullet(doc, "Inputs not available: original source code, hidden tests, internet access during inference.")
    add_bullet(doc, "Observable behavior: stdout, stderr, exit code, file effects, crashes, and timeouts.")
    add_bullet(doc, "Main failure mode: agents often start coding before they have enough behavioral evidence.")
    add_bullet(doc, "Research opportunity: design the agent-computer interface for behavioral reverse engineering, not just generic coding.")

    doc.add_heading("Why This Is More Than Probe Generation", level=1)
    doc.add_paragraph(
        "A simple pipeline that generates many tests and dumps logs to the model is unlikely to be a strong research contribution. "
        "The stronger framing is to study how a scaffold should convert black-box observations into a compact behavioral specification "
        "that helps the agent make implementation decisions under limited query, token, and repair budgets."
    )
    add_table(
        doc,
        ["Basic version", "Research-grade version"],
        [
            [
                "Run many probes.",
                "Select probes based on uncertainty and information gain: what behavior is still ambiguous?",
            ],
            [
                "Dump stdout/stderr logs.",
                "Infer a behavior map: modes, input classes, output templates, error conventions, and file effects.",
            ],
            [
                "Report mismatches.",
                "Return oracle-grounded counterexamples with inferred rules and likely implementation implications.",
            ],
            [
                "Ask the agent to fix everything.",
                "Prioritize repair targets and break repeated failure loops with new targeted probes.",
            ],
        ],
        [2.5, 4.0],
    )

    doc.add_heading("Proposed Scaffold", level=1)
    doc.add_paragraph(
        "The scaffold is a small agent-side component. It does not replace the model. It acts as a ProgramBench-specific "
        "agent-computer interface: it decides what to observe, how to summarize observations, and how to feed failures back "
        "to the coding agent."
    )
    add_table(
        doc,
        ["Component", "Role"],
        [
            ["Behavior map", "Organize observed behavior by CLI mode, input class, output template, error mode, and file effect."],
            ["Uncertainty-guided prober", "Generate the next probes to disambiguate behavior that is still unclear or under-covered."],
            ["Spec synthesizer", "Compress observations into an implementation-oriented partial specification for the agent."],
            ["Counterexample miner", "Compare reference and candidate, then attach inferred rules to representative failures."],
            ["Deadlock breaker", "Detect repeated repair loops and request new targeted observations instead of more blind edits."],
        ],
        [1.9, 4.6],
    )

    doc.add_heading("Agent Loop", level=1)
    for step in [
        "Initialize a behavior map from docs, help text, and basic executable probes.",
        "Identify uncertain behavior regions: undocumented flags, input formats, error handling, file effects, and output templates.",
        "Generate high-information probes that distinguish competing behavior hypotheses.",
        "Run the reference executable and update the behavior map with oracle-grounded observations.",
        "Synthesize a compact behavioral specification and implementation plan for the coding agent.",
        "Build and run the candidate program on the same public probes.",
        "Mine counterexamples, cluster mismatches, and guide targeted repairs until the fixed budget is exhausted.",
    ]:
        add_number(doc, step)

    doc.add_page_break()
    doc.add_heading("MVP Experiment Design", level=1)
    add_table(
        doc,
        ["Axis", "Plan"],
        [
            ["Task set", "5-10 easy ProgramBench CLI tasks; avoid synthetic sample tasks."],
            ["Primary comparison", "Baseline agent vs. the same agent plus behavioral specification scaffold."],
            ["Control", "Same model, task budget, inference setting, and no-internet constraint."],
            ["Primary metrics", "Hidden-test pass rate, full solve rate, and 95% tests-passed rate."],
            ["Secondary metrics", "Compile success, public probe match rate, regressions, reference queries, token cost, and wall-clock time."],
            ["Diagnostic metrics", "Behavior-map coverage, uncertainty reduction, counterexample reuse, and mismatch-cluster reduction."],
            ["Success signal", "Scaffold improves hidden or partial behavior metrics without simply spending much more budget."],
        ],
        [1.45, 5.05],
    )

    doc.add_heading("Ablation Plan", level=1)
    add_bullet(doc, "A: baseline mini-swe-agent.")
    add_bullet(doc, "B: baseline plus static documentation/help summary.")
    add_bullet(doc, "C: baseline plus behavior map from fixed probes.")
    add_bullet(doc, "D: baseline plus uncertainty-guided probing.")
    add_bullet(doc, "E: baseline plus oracle-grounded counterexample repair.")
    add_bullet(doc, "F: full scaffold: behavior map + uncertainty-guided probing + counterexample repair + deadlock breaker.")
    doc.add_paragraph("The first serious comparison should be A vs. F. Component ablations should follow only if the full scaffold shows a signal.")

    doc.add_heading("Positioning Against Recent Agent Work", level=1)
    add_table(
        doc,
        ["Prior direction", "Takeaway for this project"],
        [
            [
                "SWE-agent / agent-computer interfaces",
                "Agent performance depends strongly on the interface and workflow, not only the base model.",
            ],
            [
                "Agentless-style controlled pipelines",
                "A simple, inspectable scaffold can be a serious baseline if it isolates the right bottleneck.",
            ],
            [
                "Cost-aware SWE agents",
                "Exploration must be budget-aware; more probes are not automatically better.",
            ],
            [
                "Agent-generated testing studies",
                "Tests generated by agents can add cost without improving success; our probes are anchored by a real reference executable.",
            ],
            [
                "Failure-analysis work",
                "Agents often loop on the same wrong hypothesis; targeted counterexamples can break repair deadlocks.",
            ],
        ],
        [2.15, 4.35],
    )

    doc.add_page_break()
    doc.add_heading("Baseline Agents and Models", level=1)
    add_table(
        doc,
        ["Baseline", "Why use it", "Role"],
        [
            [
                "mini-swe-agent",
                "Official ProgramBench baseline interface; outputs are compatible with programbench eval.",
                "Primary harness baseline.",
            ],
            [
                "GPT-5.5",
                "OpenAI's current frontier model for complex reasoning and coding.",
                "Strong closed-model track.",
            ],
            [
                "Claude Opus 4.8 / Sonnet 4.6",
                "Current Anthropic models positioned for complex reasoning and agentic coding.",
                "Strong and balanced closed-model tracks.",
            ],
            [
                "Gemini 3.1 Pro",
                "Google DeepMind positions it as capable in agentic coding and tool use.",
                "Alternative frontier baseline.",
            ],
            [
                "GPT-5.4 mini / Claude Haiku 4.5 / Gemini Flash",
                "Cheaper or weaker models test whether scaffolding helps less capable agents more.",
                "Cost-sensitive track.",
            ],
        ],
        [1.65, 3.15, 1.7],
    )

    doc.add_heading("Current Local Evidence", level=1)
    doc.add_paragraph(
        "A local toy scaffold was migrated into this workspace and validated on /usr/bin/wc. "
        "This is not a ProgramBench score; it only confirms that the behavior-recording and mismatch-summary machinery works locally."
    )
    add_key_value_table(
        doc,
        [
            ["Probe cases", "21 generated black-box CLI cases."],
            ["Candidate", "A partial Python reimplementation of wc."],
            ["Result", "17/21 exact matches; mean partial behavior score = 0.9524."],
            ["Failure cluster", "stderr formatting for invalid/help flags."],
        ],
    )

    doc.add_heading("Infrastructure Constraint", level=1)
    doc.add_paragraph(
        "The current machine is a MacBook. Official ProgramBench containers are built for linux/amd64, so full baseline reproduction "
        "and official evaluation should run on a Linux x86-64 host. The MacBook is suitable for scaffold development, toy validation, "
        "task metadata inspection, and document/prompt design."
    )

    doc.add_heading("Next Steps", level=1)
    for step in [
        "Set up a Linux x86-64 environment with ProgramBench and mini-swe-agent.",
        "Select 5-10 easy CLI tasks and record a clean baseline run.",
        "Build the task-level behavior-map and uncertainty-guided probing runner.",
        "Run scaffold-assisted agents under the same budget as the baseline.",
        "Compare hidden scores, partial behavior scores, behavior-map coverage, cost, and qualitative trajectories.",
    ]:
        add_bullet(doc, step)

    doc.add_heading("Sources", level=1)
    for source in [
        "ProgramBench GitHub: https://github.com/facebookresearch/programbench",
        "ProgramBench paper: https://arxiv.org/abs/2605.03546",
        "SWE-agent paper: https://arxiv.org/abs/2405.15793",
        "Agentless paper: https://arxiv.org/abs/2407.01489",
        "mini-swe-agent ProgramBench docs: https://mini-swe-agent.com/latest/usage/programbench/",
        "OpenAI model docs: https://developers.openai.com/api/docs/models",
        "Anthropic Claude model docs: https://docs.anthropic.com/en/docs/about-claude/models/overview",
        "Google DeepMind Gemini Pro docs: https://deepmind.google/models/gemini/pro/",
    ]:
        add_bullet(doc, source)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    build_doc()
